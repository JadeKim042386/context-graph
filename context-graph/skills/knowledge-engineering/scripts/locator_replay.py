"""Bounded local locator replay. Location validity is not semantic evidence.

Optional binary parsers execute in a time-limited subprocess. No downloads,
OCR/transcription, external references, parser installation or content promotion.
"""
import csv
import base64
import io
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path

MAX_INPUT = 262144
STATES = {'valid','unsupported','unverified','unlocatable','oversize'}


class Anchors(HTMLParser):
    def __init__(self, target):
        super().__init__(convert_charrefs=True)
        self.target, self.stack, self.matches = target, [], []
    def handle_starttag(self, tag, attrs):
        attrs=dict(attrs)
        if tag in {'br','img','meta','link','input','hr','area','base','embed','source','wbr','param','col','track'}:
            if attrs.get('id') == self.target: self.matches.append([])
            return
        match=[] if attrs.get('id')==self.target or (tag=='a' and attrs.get('name')==self.target) else None
        if match is not None: self.matches.append(match)
        self.stack.append((tag,match))
    def handle_endtag(self,tag):
        if self.stack and self.stack[-1][0] == tag: self.stack.pop()
        elif any(t==tag for t,_ in self.stack): raise ValueError('ambiguous_html')
    def handle_data(self,data):
        for _,match in self.stack:
            if match is not None: match.append(data)


def region(value,width,height):
    points=[float(x) for x in value.split(',')]
    if len(points)!=4 or not all(math.isfinite(x) for x in points): return False
    x0,y0,x1,y1=points
    return 0 <= x0 < x1 <= width and 0 <= y0 < y1 <= height


def _replay(body, locator, quote):
    try:
        if not isinstance(locator,str) or len(locator)>512 or (quote is not None and (not isinstance(quote,str) or len(quote)>8192)): return 'unverified'
        m=re.fullmatch(r'line:([1-9][0-9]*)(?:-([1-9][0-9]*))?',locator)
        text=None
        if m:
            lines=body.decode('utf-8').splitlines(); start,end=int(m[1]),int(m[2] or m[1])
            if not start<=end<=len(lines): return 'unlocatable'
            text='\n'.join(lines[start-1:end])
        elif locator.startswith('html:'):
            m=re.fullmatch(r'html:#([^\s#]+)',locator)
            if not m: return 'unverified'
            parser=Anchors(m[1]); parser.feed(body.decode('utf-8')); parser.close()
            if len(parser.matches)!=1 or any(x is not None for _,x in parser.stack): return 'unlocatable'
            text=''.join(parser.matches[0])
        elif locator.startswith('pdf:'):
            from pypdf import PdfReader
            m=re.fullmatch(r'pdf:page=([1-9][0-9]*)(?:;bbox=([0-9.,]+))?',locator)
            if not m: return 'unverified'
            pdf=PdfReader(io.BytesIO(body),strict=True)
            if pdf.is_encrypted or len(pdf.pages)>1000 or int(m[1])>len(pdf.pages): return 'unlocatable'
            page=pdf.pages[int(m[1])-1]
            if m[2]:
                if not region(m[2],float(page.mediabox.width),float(page.mediabox.height)): return 'unlocatable'
                # pypdf cannot guarantee glyph containment for arbitrary transformed content.
                if quote: return 'unsupported'
            text=page.extract_text() if quote else ''
        elif locator.startswith('table:'):
            m=re.fullmatch(r'table:sheet=([^;]{1,31});cell=([A-Z]{1,3}[1-9][0-9]{0,6}(?::[A-Z]{1,3}[1-9][0-9]{0,6})?)',locator)
            if not m: return 'unverified'
            if not body.startswith(b'PK') and m[1]=='CSV':
                def cell(value):
                    match=re.fullmatch(r'([A-Z]+)([0-9]+)',value)
                    col=0
                    for c in match[1]: col=col*26+ord(c)-64
                    return col,int(match[2])
                bounds=m[2].split(':');c0,r0=cell(bounds[0]);c1,r1=cell(bounds[-1])
                rows=list(csv.reader(io.StringIO(body.decode('utf-8-sig')),strict=True))
                if not 1<=r0<=r1<=len(rows) or not 1<=c0<=c1 or (r1-r0+1)*(c1-c0+1)>10000: return 'unlocatable'
                if any(len(row)<c1 for row in rows[r0-1:r1]):return 'unlocatable'
                selected=[row[c0-1:c1] for row in rows[r0-1:r1]]
                if any(value.startswith('=') for row in selected for value in row):return 'unsupported'
                text='\n'.join('\t'.join(row) for row in selected)
                return 'valid' if not quote or quote in text else 'unlocatable'
            import openpyxl
            from openpyxl.utils.cell import range_boundaries
            with zipfile.ZipFile(io.BytesIO(body)) as z:
                if sum(x.file_size for x in z.infolist())>16*1024*1024 or len(z.infolist())>1000: return 'unlocatable'
            book=openpyxl.load_workbook(io.BytesIO(body),read_only=True,data_only=False,keep_links=False)
            try:
                if m[1] not in book.sheetnames: return 'unlocatable'
                sheet=book[m[1]]; c0,r0,c1,r1=range_boundaries(m[2])
                if not 1<=r0<=r1<=sheet.max_row or not 1<=c0<=c1<=sheet.max_column or (r1-r0+1)*(c1-c0+1)>10000: return 'unlocatable'
                values=list(sheet.iter_rows(min_row=r0,max_row=r1,min_col=c0,max_col=c1))
                # No formula evaluation or cached-value correctness claim.
                if any(c.data_type=='f' for row in values for c in row): return 'unsupported'
                text='\n'.join('\t'.join('' if c.value is None else str(c.value) for c in row) for row in values)
            finally: book.close()
        elif locator.startswith('image:'):
            from PIL import Image
            m=re.fullmatch(r'image:bbox=([0-9.,]+)',locator)
            if not m: return 'unverified'
            with Image.open(io.BytesIO(body)) as image:
                if image.width*image.height>20000000: return 'unlocatable'
                if not region(m[1],image.width,image.height): return 'unlocatable'
                image.verify()
            return 'unsupported' if quote else 'valid'
        elif locator.startswith('av:'):
            m=re.fullmatch(r'av:t=([0-9]{2}:[0-5][0-9]:[0-5][0-9](?:\.[0-9]+)?)-([0-9]{2}:[0-5][0-9]:[0-5][0-9](?:\.[0-9]+)?)',locator)
            if not m: return 'unverified'
            probe=shutil.which('ffprobe')
            if not probe: return 'unsupported'
            if not (body.startswith((b'RIFF',b'fLaC',b'ID3',b'\xff\xfb',b'\xff\xf3',b'\xff\xf2')) or body[4:8]==b'ftyp'): return 'unlocatable'
            with tempfile.TemporaryDirectory(prefix='locator-') as tmp:
                path=Path(tmp)/'media.bin'; path.write_bytes(body)
                command=[probe,'-v','error','-protocol_whitelist','file','-format_whitelist','wav,flac,mp3,mov', '-show_entries','format=duration','-of','json',str(path)]
                result=subprocess.run(command,capture_output=True,timeout=2)
            if result.returncode or len(result.stdout)>16384: return 'unlocatable'
            duration=float(json.loads(result.stdout)['format']['duration'])
            def seconds(value):
                h,minutes,s=value.split(':'); return int(h)*3600+int(minutes)*60+float(s)
            if not math.isfinite(duration) or not 0<=seconds(m[1])<seconds(m[2])<=duration: return 'unlocatable'
            return 'unsupported' if quote else 'valid'
        else: return 'unsupported'
        return 'valid' if not quote or quote in text else 'unlocatable'
    except ImportError: return 'unsupported'
    except Exception: return 'unlocatable'


def replay(body, locator, quote=None):
    if len(body)>MAX_INPUT: return 'oversize'
    if isinstance(locator,str) and locator.startswith(('line:','html:')):
        return _replay(body,locator,quote)
    try:
        raw=json.dumps({'body':base64.b64encode(body).decode(),'locator':locator,'quote':quote}).encode()
        result=subprocess.run([sys.executable,'-B',str(Path(__file__).resolve()),'--worker'],input=raw,capture_output=True,timeout=4)
        answer=json.loads(result.stdout) if result.returncode==0 and len(result.stdout)<1024 else None
        return answer if answer in STATES else 'unlocatable'
    except (OSError,ValueError,subprocess.TimeoutExpired): return 'unlocatable'


if __name__=='__main__':
    import resource
    resource.setrlimit(resource.RLIMIT_CPU,(3,3))
    resource.setrlimit(resource.RLIMIT_FSIZE,(1024*1024,1024*1024))
    try:
        data=json.loads(sys.stdin.buffer.read(400000))
        body=base64.b64decode(data['body'],validate=True)
        state=_replay(body,data['locator'],data.get('quote')) if len(body)<=MAX_INPUT else 'oversize'
    except Exception: state='unlocatable'
    print(json.dumps(state))
